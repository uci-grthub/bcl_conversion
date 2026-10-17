"""Cross-sheet checks in src/metadata_validation.py.

Covers the workbook mistakes these checks exist to catch: a Barcode List
project name that differs from the Summary, a project name pasted into the
LabID column, sample names with spaces, and short single indexes sharing a
lane with longer dual indexes.
"""
import os

import openpyxl
import pandas as pd

from metadata_validation import (
    check_index_collisions,
    check_lane_group_consistency,
    check_name_safety,
    validate_metadata_and_write_report,
)


def _messages(issues):
    return [it["message"] for it in issues]


SUMMARY = pd.DataFrame([
    {"Lane": 7, "Gr": 2, "Order ID": "1026I-26", "Lab ID": "LinE",
     "Project Name": "LinE_RNAseq_PoolA_8plex", "Contact (E-mail)": "nguyenq@aurekabio.com"},
    {"Lane": 7, "Gr": 9, "Order ID": "0926I-36", "Lab ID": "EmerJ",
     "Project Name": "EmerJ_ISO1 single female", "Contact (E-mail)": "alyssenw@uci.edu"},
])


def _barcode_list(rows):
    return pd.DataFrame(rows, columns=[
        "Lane", "Group", "Order ID", "LabID", "Contact", "Project name"])


# --- check_lane_group_consistency -------------------------------------------

def test_project_typo_on_barcode_list_is_flagged():
    bl = _barcode_list([
        (7, 2, "1026I-26", "LinE", "nguyenq@aurekabio.com", "LinE_RNAsq_PoolA_8plex"),
        (7, 2, "1026I-26", "LinE", "nguyenq@aurekabio.com", "LinE_RNAsq_PoolA_8plex"),
    ])
    msgs = _messages(check_lane_group_consistency({"Summary": SUMMARY, "Barcode List": bl}))
    assert msgs == ["Project name 'LinE_RNAsq_PoolA_8plex' does not match Summary "
                    "'LinE_RNAseq_PoolA_8plex' for lane 7 group 2"]


def test_project_name_in_labid_column_is_flagged():
    bl = _barcode_list([
        (7, 9, "0926I-36", "EmerJ_ISO1 single female", "alyssenw@uci.edu",
         "EmerJ_ISO1 single female"),
    ])
    msgs = _messages(check_lane_group_consistency({"Summary": SUMMARY, "Barcode List": bl}))
    assert msgs == ["Lab ID 'EmerJ_ISO1 single female' does not match Summary 'EmerJ' "
                    "for lane 7 group 9"]


def test_spelling_differences_the_pipeline_normalizes_are_not_flagged():
    # Spaces vs underscores in a project name and e-mail case are the same value.
    bl = _barcode_list([
        (7, 9, "0926I-36", "EmerJ", "AlysSenW@uci.edu", "EmerJ_ISO1_single_female"),
    ])
    assert check_lane_group_consistency({"Summary": SUMMARY, "Barcode List": bl}) == []


def test_no_summary_means_nothing_to_compare():
    bl = _barcode_list([(1, 1, "X", "Y", "z@x.org", "P")])
    assert check_lane_group_consistency({"Barcode List": bl}) == []


# --- check_name_safety -------------------------------------------------------

def test_renamed_samples_are_reported_once_per_project():
    df = pd.DataFrame({
        "Lane": [7] * 5,
        "Project name": ["MurnJ_SL"] * 5,
        "Sample Name": ["Input C1", "Input C2", "Input C3", "Input T4", "OK_name"],
    })
    msgs = _messages(check_name_safety("Barcode List", df, examples=2))
    assert msgs == ["4 sample names in MurnJ_SL will be renamed for DRAGEN: "
                    "'Input C1' -> 'Input_C1', 'Input C2' -> 'Input_C2' (+2 more)"]


def test_doubled_underscores_count_as_a_rename():
    df = pd.DataFrame({"Lane": [6], "Project name": ["ChesN"],
                       "Sample Name": ["INT_TACplusHFD__2L__LV"]})
    msgs = _messages(check_name_safety("Barcode List", df))
    assert msgs == ["1 sample name in ChesN will be renamed for DRAGEN: "
                    "'INT_TACplusHFD__2L__LV' -> 'INT_TACplusHFD_2L_LV'"]


def test_names_that_collapse_together_are_flagged():
    df = pd.DataFrame({
        "Lane": [1, 1, 1, 2],
        "Project name": ["P"] * 4,
        "Sample Name": ["A_B", "A B", "A B", "A-B"],
    })
    msgs = _messages(check_name_safety("Barcode List", df))
    assert "Sample names 'A_B' and 'A B' both become 'A_B' after DRAGEN-safe renaming" in msgs
    # Reported once, not once per repeated row.
    assert sum("both become" in m for m in msgs) == 1


def test_project_name_with_spaces_is_flagged():
    df = pd.DataFrame({"Lane": [7], "Project Name": ["EmerJ_ISO1 single female"]})
    msgs = _messages(check_name_safety("Summary", df))
    assert msgs == ["Project name 'EmerJ_ISO1 single female' will be delivered as "
                    "'EmerJ_ISO1_single_female'"]


# --- check_index_collisions --------------------------------------------------

def _bl(rows):
    return pd.DataFrame(rows, columns=[
        "Lane", "Group", "Project", "Sample Name", "i7 Barcode Sequence", "i5 Barcode Sequence"])


def test_short_single_index_near_longer_dual_index_is_one_issue_per_project_pair():
    df = _bl([
        (7, 7, "McClM_Chicken", "S1", "ATAGAGAG", "TAAGGCGA"),
        (7, 7, "McClM_Chicken", "S2", "ATAGAGAG", "CGTACTAG"),
        (7, 12, "ShiY_PASseq", "B2", "GTAGAG", None),
    ])
    issues = check_index_collisions({"Barcode List": df})
    assert len(issues) == 1
    msg = issues[0]["message"]
    assert msg.startswith("Lane 7: indexes in projects McClM_Chicken and ShiY_PASseq")
    assert "i7 distance 1, no i5 to separate them (2 close pairs in all)" in msg
    assert issues[0]["lane"] == "7"


def test_prefix_collision_is_unresolvable_and_routes_short_project():
    df = _bl([
        (7, 7, "LongP", "L1", "GTAGAGGA", None),
        (7, 12, "ShortP", "S1", "GTAGAG", None),
    ])
    msgs = _messages(check_index_collisions({"Barcode List": df}))
    assert len(msgs) == 1
    assert msgs[0].startswith("Unresolvable index collision: lane 7:")
    assert "demultiplex project 'ShortP' from Undetermined with fqtk" in msgs[0]


def test_distinct_i5_separates_close_i7():
    df = _bl([
        (3, 1, "A", "a1", "ACGTACGT", "AAAAAAAA"),
        (3, 2, "B", "b1", "ACGTACGA", "CCCCCCCC"),
    ])
    assert check_index_collisions({"Barcode List": df}) == []


def test_other_lanes_and_post_hoc_projects_are_ignored():
    df = _bl([
        (3, 1, "A", "a1", "GTAGAGGA", None),
        (4, 1, "B", "b1", "GTAGAG", None),
        (3, 2, "C_fqtk", "c1", "GTAGAG", None),
    ])
    assert check_index_collisions({"Barcode List": df}) == []


def test_summary_rows_are_not_samples():
    summary = pd.DataFrame({"Lane": [1, 1], "Project": ["A", "B"],
                            "index": ["ACGTACGT", "ACGTACGT"]})
    assert check_index_collisions({"Summary": summary}) == []


# --- end to end --------------------------------------------------------------

def _write_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([None, None, None, "Flowcell:TEST"])
    ws.append([])
    ws.append(["Lane", "Gr", "Order ID", "Lab ID", "Project Name", "Contact (E-mail)",
               "Sample sheet tab", "Masking"])
    ws.append([1, 1, "1026I-05", "KencS", "KencS_10x3V4_X1", "k@ucr.edu",
               "10x Dual Index", "R1:28, I1:10, I2:10, R2:90"])
    ws.append([2, 1, "1026I-26", "LinE", "LinE_RNAseq_PoolA_8plex", "n@aurekabio.com",
               "Barcode List", "R1:151, I1:10, I2:10, R2:151"])

    ws = wb.create_sheet("10x Dual Index")
    ws.append(["[Data]"])
    ws.append(["Lane", "Sample_ID", "Sample_Name", "index", "index2", "Sample_Project",
               None, "Lane", "Group", "Lab ID", "Email"])
    ws.append([1, "SI-TT-D5", "W8", "TGGTTCGGGT", "CTCCTGCCAC", None,
               None, 1, 1, "KencS", "k@ucr.edu"])

    ws = wb.create_sheet("Barcode List")
    ws.append([])
    ws.append(["Lane", "Group", "Order ID", "LabID", "Contact", "Project name",
               "Sample Name", "Index Name", "i7 Barcode Sequence", "i5 Barcode Sequence"])
    ws.append([2, 1, "1026I-26", "LinE", "n@aurekabio.com", "LinE_RNAsq_PoolA_8plex",
               "XIA002_TR_D0_4", "UDP0257", "GTAACAATCT", "GTATACAGAG"])
    wb.save(path)


def test_report_flags_typo_and_not_10x_projects(tmp_path):
    os.chdir(tmp_path)
    _write_workbook("in.xlsx")
    validate_metadata_and_write_report("in.xlsx", out_xlsx="out.xlsx")
    msgs = pd.read_excel("out.xlsx", sheet_name="RECOMMENDED_CHANGES")["message"].tolist()

    assert ("Project name 'LinE_RNAsq_PoolA_8plex' does not match Summary "
            "'LinE_RNAseq_PoolA_8plex' for lane 2 group 1") in msgs
    # KencS lives on the 10x tab its Summary row names, so Barcode List is not expected to list it.
    assert not any("KencS" in m and "missing from Barcode List" in m for m in msgs)
